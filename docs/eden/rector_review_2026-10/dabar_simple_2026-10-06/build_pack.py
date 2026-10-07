"""Build compact accepted Word and HTML policy documents from the retained page sources."""
from pathlib import Path
from datetime import datetime, timezone
import re
import html
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent
LINK = re.compile(r'\[([^\]]+)\]\(([^\s)]+)\)')

def add_inline(p, text):
    at = 0
    for m in LINK.finditer(text):
        p.add_run(text[at:m.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(m[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '185A82'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props)
        t = OxmlElement('w:t'); t.text = m[1]; run.append(t)
        link.append(run); p._p.append(link)
        at = m.end()
    p.add_run(text[at:])

def inline_html(text):
    def replace(m):
        target = '04_Contact.html' if m[2] == 'https://repository.uist.edu.mk/info/contact' else m[2]
        return f'<a href="{html.escape(target, quote=True)}">{html.escape(m[1])}</a>'
    return LINK.sub(replace, html.escape(text))

def build(source):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(18); sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(22); sec.right_margin = Mm(22)
    sec.footer_distance = Mm(9)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]; style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    normal = doc.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    title = doc.styles['Title']; title.font.size = Pt(22); title.font.bold = True
    title.paragraph_format.space_after = Pt(10); title.paragraph_format.keep_with_next = True
    heading = doc.styles['Heading 1']; heading.font.size = Pt(12); heading.font.bold = True
    heading.paragraph_format.space_before = Pt(10); heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.add_run('UIST Digital Repository').font.size = Pt(8)
    foot = sec.footer.paragraphs[0]
    foot.add_run('Accepted by rector   |   Version 6 October 2026   |   ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    body = []
    for line in source.read_text().splitlines():
        line = line.strip()
        if not line: continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title'); body.append(f'<h1>{html.escape(line[2:])}</h1>')
        elif line.startswith('## '):
            if source.stem == '01_Repository_Policies' and line == '## Metadata reuse and curation':
                doc.add_page_break()
            anchor = re.sub(r'[^a-z0-9]+', '-', line[3:].lower()).strip('-')
            doc.add_paragraph(line[3:], 'Heading 1'); body.append(f'<h2 id="{anchor}">{html.escape(line[3:])}</h2>')
        else:
            add_inline(doc.add_paragraph(), line); body.append(f'<p>{inline_html(line)}</p>')
    title_text = source.read_text().splitlines()[0][2:]
    doc.core_properties.title = title_text
    doc.core_properties.author = 'UIST Digital Repository'
    doc.core_properties.subject = 'Simplified public policy accepted by the rector; acceptance reported 6 October 2026'
    doc.core_properties.created = datetime(2026, 10, 6, 10, 0, tzinfo=timezone.utc)
    doc.save(BASE / (source.stem + '.docx'))
    nav = ' | '.join(f'<a href="{stem}.html">{label}</a>' for stem, label in [
        ('01_Repository_Policies', 'Repository policies'), ('02_Accessibility_Statement', 'Accessibility statement'),
        ('03_Privacy_Notice', 'Privacy notice'), ('04_Contact', 'Contact')])
    page = '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
    page += f'<title>{html.escape(title_text)}</title><style>body{{font:18px/1.55 system-ui,sans-serif;color:#171717;margin:0 auto;padding:2rem 1.5rem;max-width:760px}}h1{{font-size:2rem;line-height:1.2}}h2{{font-size:1.15rem;margin-top:1.7rem}}a{{color:#185a82}}a:focus-visible{{outline:3px solid #171717;outline-offset:3px}}footer{{margin-top:2.5rem;font-size:.9rem}}.status{{font-size:.85rem;color:#555}}</style>'
    page += '<p class="status">Accepted by rector · Version 6 October 2026</p><main>' + '\n'.join(body) + '</main><footer><nav aria-label="Repository information">' + nav + '</nav></footer></html>'
    (BASE / 'html').mkdir(exist_ok=True)
    (BASE / 'html' / (source.stem + '.html')).write_text(page)
    text = LINK.sub(lambda m:m[1], source.read_text())
    print(source.stem, len(text.split()), 'words')

if __name__ == '__main__':
    for source in sorted((BASE/'pages').glob('*.md')): build(source)
