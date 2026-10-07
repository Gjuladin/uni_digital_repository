"""Build the six editable rector-review documents from their retained Markdown sources."""
from pathlib import Path
import re
import sys
from datetime import datetime, timezone
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent


def inline(p, text):
    """Preserve real source hyperlinks, without importing Markdown punctuation."""
    pattern = re.compile(r'\[([^\]]+)\]\((https?://[^\s)]+)\)')
    at = 0
    for match in pattern.finditer(text):
        p.add_run(text[at:match.start()])
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), p.part.relate_to(match.group(2), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '1A4A72'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props)
        t = OxmlElement('w:t'); t.text = match.group(1); run.append(t)
        hyperlink.append(run); p._p.append(hyperlink)
        at = match.end()
    p.add_run(text[at:])


def table(doc, rows):
    tab = doc.add_table(rows=1, cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    widths = [46, 84, 40] if len(rows[0]) == 3 else [62, 108]
    for col, width in zip(tab.columns, widths): col.width = Mm(width)
    pr = tab._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        el = OxmlElement('w:' + edge); el.set(qn('w:w'), '95'); el.set(qn('w:type'), 'dxa'); margins.append(el)
    pr.append(margins)
    header = OxmlElement('w:tblHeader'); tab.rows[0]._tr.get_or_add_trPr().append(header)
    for i, texts in enumerate(rows):
        row = tab.rows[0] if i == 0 else tab.add_row()
        # Keep an individual row together; let the table continue with a repeated header.
        cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
        for j, text in enumerate(texts):
            cell = row.cells[j]; cell.width = Mm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            inline(p, text)
            for r in p.runs: r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0, 0, 0)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E3EAF0'); cell._tc.get_or_add_tcPr().append(shade)
                for r in p.runs: r.bold = True
            elif i % 2 == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F7F8FA'); cell._tc.get_or_add_tcPr().append(shade)
    spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(1); spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.4; spacer.add_run().font.size = Pt(3)


def build(source):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(19); sec.bottom_margin = Mm(19)
    sec.left_margin = Mm(20); sec.right_margin = Mm(20)
    sec.footer_distance = Mm(8)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number', 'List Bullet']:
        style = doc.styles[name]; style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
    # The bundled default template carries a blue Title border; remove it at source.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal = doc.styles['Normal']; normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.08; normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    title = doc.styles['Title']; title.font.size = Pt(21); title.font.bold = True
    title.paragraph_format.space_after = Pt(8); title.paragraph_format.keep_with_next = True
    for name, size in [('Heading 1', 13), ('Heading 2', 11.5)]:
        style = doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
        style.paragraph_format.space_before = Pt(11); style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    # Visible draft status is needed to prevent proposed rules being mistaken for adopted policy.
    foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.paragraph_format.space_after = Pt(0)
    r = foot.add_run('UIST   |   Draft for review   |   '); r.font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in cells): rows.append(cells)
                i += 1
            table(doc, rows); continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        elif re.match(r'^\d+\. ', line):
            p = doc.add_paragraph(style='List Number'); inline(p, re.sub(r'^\d+\. ', '', line))
        else:
            p = doc.add_paragraph(); inline(p, line)
            if i == 2: p.paragraph_format.space_after = Pt(12); [setattr(r.font, 'size', Pt(9.5)) for r in p.runs]
            if line.startswith('['):
                p.paragraph_format.space_after = Pt(5)
                for r in p.runs: r.font.size = Pt(9)
        i += 1
    doc.core_properties.title = lines[0][2:]
    doc.core_properties.subject = 'UIST repository policies and registry readiness'
    doc.core_properties.author = 'UIST Digital Repository development team'
    doc.core_properties.keywords = 'Draft; rector review; UIST; re3data; FAIRsharing'
    doc.core_properties.created = datetime(2026, 10, 3, 10, 0, tzinfo=timezone.utc)
    dest = BASE / (source.stem + '.docx'); doc.save(dest)
    return dest


if __name__ == '__main__':
    sources = [BASE / 'drafts' / (stem + '.md') for stem in sys.argv[1:]] if len(sys.argv) > 1 else sorted((BASE / 'drafts').glob('*.md'))
    for source in sources:
        output = build(source)
        print(output.name)
